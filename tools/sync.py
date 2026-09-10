#!/usr/bin/env python3
"""Pull the roadmap workbook from its share link and rebuild the chart.

    python3 tools/sync.py                      # uses the link below
    python3 tools/sync.py <share-or-file-url>  # or any other link

Run this from a machine that can reach the share -- a laptop on the SURF or
university network. GitHub's runners cannot: surf.works.surf.nl does not
accept connections from them, so the scheduled workflow is switched off.

Then commit what changed:

    git add data gantt.html && git commit -m "Update roadmap" && git push
"""
import ssl
import sys
import urllib.error
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "data" / "durf-roadmap.xlsx"
DEFAULT_URL = "https://surf.works.surf.nl/s/NMABkFM9pMrdkBJ"
TIMEOUT = 60


def download(url):
    """Fetch the bytes behind a share link, following Nextcloud's convention."""
    url = url.rstrip("/")
    if not url.endswith("/download"):
        # A Nextcloud/ownCloud /s/<token> link serves a viewer page; the file
        # itself is one segment further on.
        url += "/download"
    print(f"fetching {url}")
    request = urllib.request.Request(url, headers={"User-Agent": "durf-gantt sync"})
    try:
        with urllib.request.urlopen(request, timeout=TIMEOUT,
                                    context=ssl.create_default_context()) as res:
            return res.read()
    except urllib.error.HTTPError as err:
        raise SystemExit(f"error: the server answered {err.code} {err.reason}")
    except (urllib.error.URLError, TimeoutError) as err:
        raise SystemExit(
            f"error: could not reach the server ({err}).\n"
            "       Are you on a network that can open the share link?")


def check(payload):
    """A password prompt, an expiry or a spent download limit all arrive as
    HTML with a 200, so never trust the status code on its own."""
    if payload[:4] != b"PK\x03\x04":
        head = payload[:200].decode("utf-8", "replace").strip()
        raise SystemExit(
            "error: that link did not return a spreadsheet.\n"
            f"       It starts with: {head[:120]}\n"
            "       Check the share is public, unexpired, has no password and\n"
            "       no download limit.")

    from openpyxl import load_workbook  # imported late so the error above is cheap
    import io
    workbook = load_workbook(io.BytesIO(payload), read_only=True)
    missing = {"Themes", "Activities"} - set(workbook.sheetnames)
    if missing:
        raise SystemExit(f"error: the workbook has no {', '.join(sorted(missing))} sheet")
    print("  sheets: " + ", ".join(workbook.sheetnames))


def main():
    url = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_URL
    payload = download(url)
    check(payload)

    before = TARGET.read_bytes() if TARGET.exists() else b""
    if payload == before:
        print("Already up to date — nothing changed.")
        return

    TARGET.write_bytes(payload)
    print(f"  wrote {TARGET.relative_to(ROOT)} ({len(payload) / 1024:.1f} kB)")

    sys.path.insert(0, str(ROOT / "tools"))
    import build
    build.main()
    print("\nDone. Commit the result:\n"
          "    git add data gantt.html && git commit -m \"Update roadmap\" && git push")


if __name__ == "__main__":
    main()
