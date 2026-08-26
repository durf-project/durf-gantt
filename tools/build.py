#!/usr/bin/env python3
"""Bake data/durf-roadmap.xlsx into index.html and export per-sheet CSVs.

    python3 tools/build.py

index.html keeps a block delimited by DURF-DATA markers. This script replaces the
JSON inside it with the workbook's raw sheet rows. The page normalises those rows
in JavaScript, using exactly the same code path as a workbook a user drops into
the page at runtime -- so a baked page and a freshly-loaded file can never drift.
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "durf-roadmap.xlsx"
HTML = ROOT / "index.html"
CSV_DIR = ROOT / "data" / "csv"

SKIP_SHEETS = {"READ ME"}
BEGIN = "/* == DURF-DATA:BEGIN == */"
END = "/* == DURF-DATA:END == */"


def cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    return str(value).strip()


def read_sheets(path):
    wb = load_workbook(path, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        rows = []
        for raw in ws.iter_rows(values_only=True):
            row = [cell(v) for v in raw]
            while row and row[-1] == "":
                row.pop()
            rows.append(row)
        while rows and not any(str(v) for v in rows[-1]):
            rows.pop()
        if rows:
            sheets[ws.title] = rows
    return sheets


def write_csvs(sheets):
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    for name, rows in sheets.items():
        width = max(len(r) for r in rows)
        out = CSV_DIR / (name.lower().replace(" ", "-") + ".csv")
        with out.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            for row in rows:
                w.writerow(list(row) + [""] * (width - len(row)))
        print(f"  {out.relative_to(ROOT)}")


def bake(sheets):
    html = HTML.read_text(encoding="utf-8")
    if BEGIN not in html or END not in html:
        sys.exit(f"marker block not found in {HTML.name}")
    payload = json.dumps(sheets, ensure_ascii=False, separators=(",", ":"))
    # </script> inside data would close the tag early.
    payload = payload.replace("</", "<\\/")
    block = f"{BEGIN}\nwindow.__DURF_SHEETS__ = {payload};\n{END}"
    html = re.sub(re.escape(BEGIN) + r".*?" + re.escape(END), lambda _: block,
                  html, flags=re.S)
    HTML.write_text(html, encoding="utf-8")
    print(f"  {HTML.name}  ({len(payload) / 1024:.1f} kB of data)")


def main():
    if not XLSX.exists():
        sys.exit(f"missing {XLSX} -- run tools/seed_workbook.py first")
    sheets = read_sheets(XLSX)
    print(f"read {XLSX.relative_to(ROOT)}: "
          + ", ".join(f"{k} ({len(v) - 1})" for k, v in sheets.items()))
    write_csvs(sheets)
    if HTML.exists():
        bake(sheets)
    else:
        print(f"  {HTML.name} not found yet -- skipped baking")


if __name__ == "__main__":
    main()
