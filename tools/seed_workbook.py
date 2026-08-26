#!/usr/bin/env python3
"""Generate data/durf-roadmap.xlsx from the transcribed deck content in seed_data.py.

Run once to create the workbook. After that the workbook is the source of truth --
re-running this overwrites any edits.

    python3 tools/seed_workbook.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import seed_data as seed

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "durf-roadmap.xlsx"

INK = "FF050E1D"          # --sds--color--black
GREY_100 = "FFF4F6F8"     # --sds--color--gray--100
GREY_300 = "FFB2B6BE"     # --sds--color--gray--300
BLUE_400 = "FF0077C8"     # --sds--color--blue--400

HEAD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BLUE_400)
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="FF5E6873")
THIN = Side(style="thin", color=GREY_300)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TAGS = ["report", "standard", "signed agreement", "legal", "configuration",
        "software", "service", "training", "event", "feature", "data", "sla"]
STATUSES = ["Not started", "In progress", "Done", "At risk", "Cancelled"]


def style_header(ws, ncols, freeze="A2"):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def body(ws, first_row, ncols, wrap_cols=()):
    for row in ws.iter_rows(min_row=first_row, max_row=ws.max_row, max_col=ncols):
        for cell in row:
            cell.border = BOX
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(
                vertical="top", wrap_text=cell.column in wrap_cols)


def sheet_readme(wb):
    ws = wb.create_sheet("READ ME", 0)
    ws.sheet_properties.tabColor = "0077C8"
    set_widths(ws, [4, 108])
    lines = [
        ("h1", "DURF roadmap — how to edit this workbook"),
        ("p", ""),
        ("p", "This workbook fills the DURF Gantt chart (index.html). Edit here, "
              "never in the HTML."),
        ("p", ""),
        ("h2", "Two ways to get your edits into the chart"),
        ("li", "Quick preview — open index.html in a browser, click “Load data”, "
               "and pick this .xlsx file. Nothing is uploaded; the file is read "
               "in your browser."),
        ("li", "Publish — ask a developer to run  python3 tools/build.py  once. "
               "That bakes this workbook into index.html so the published page "
               "needs no extra files."),
        ("p", ""),
        ("h2", "The sheets"),
        ("li", "Themes — the six DURF themes. One row per theme. The colour column "
               "drives the chart colours."),
        ("li", "Activities — the Gantt bars. One row per activity. This is the "
               "sheet you will edit most."),
        ("li", "Budget — budget lines shown when a theme is opened."),
        ("li", "Open questions — the “what about …?” notes from the deck."),
        ("li", "Settings — title, subtitle, project start month, project length."),
        ("p", ""),
        ("h2", "Rules that matter"),
        ("li", "Do not rename the sheets or the header row — the chart looks them "
               "up by name."),
        ("li", "start_month / end_month are project months counted from month 1, "
               "not calendar dates. Month 1 is set in Settings."),
        ("li", "For an activity that is a single moment, put the same number in "
               "start_month and end_month."),
        ("li", "milestone_months is a comma-separated list, e.g. 12, 24, 36, 48. "
               "Each one becomes a diamond on the bar. Leave blank if there are "
               "none."),
        ("li", "outcomes: separate several outcomes with a vertical bar |"),
        ("li", "tags: comma-separated. They become the filter chips above the "
               "chart. Reuse existing spellings so filters stay tidy."),
        ("li", "links: write them as   Label :: https://url   and separate "
               "several with a vertical bar |"),
        ("li", "Colours are hex codes like #EE7628. Text colour on the bars is "
               "picked automatically for contrast."),
        ("li", "Rows are drawn in sheet order. Sort or move rows to reorder the "
               "chart."),
        ("li", "Delete a row to remove a bar. Add a row at the bottom to add one."),
        ("p", ""),
        ("h2", "If you prefer CSV"),
        ("li", "Save any single sheet as CSV and load that instead. The chart "
               "recognises a sheet by its header row, so Activities.csv, "
               "Themes.csv and Budget.csv all work on their own."),
    ]
    r = 2
    for kind, text in lines:
        cell = ws.cell(row=r, column=2, value=text)
        if kind == "h1":
            cell.font = Font(name="Calibri", size=16, bold=True, color=INK)
            ws.row_dimensions[r].height = 24
        elif kind == "h2":
            cell.font = Font(name="Calibri", size=12, bold=True, color=BLUE_400)
        elif kind == "li":
            cell.value = "•  " + text
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 15 * (len(text) // 95 + 1))
        else:
            cell.font = Font(name="Calibri", size=11)
        r += 1
    ws.sheet_view.showGridLines = False


def sheet_settings(wb):
    ws = wb.create_sheet("Settings")
    ws.append(["key", "value", "what it does"])
    for k, v, note in seed.SETTINGS:
        ws.append([k, v, note])
    style_header(ws, 3)
    set_widths(ws, [22, 62, 58])
    body(ws, 2, 3, wrap_cols=(2, 3))
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).font = NOTE_FONT


def sheet_themes(wb):
    ws = wb.create_sheet("Themes")
    ws.append(["theme_no", "name", "short_name", "colour", "goal", "lead",
               "partners"])
    for t in seed.THEMES:
        ws.append([t["no"], t["name"], t["short"], t["color"], t["goal"],
                   t["lead"], t["partners"]])
    style_header(ws, 7, freeze="B2")
    set_widths(ws, [9, 46, 22, 10, 70, 16, 40])
    body(ws, 2, 7, wrap_cols=(2, 5, 7))
    for r, t in enumerate(seed.THEMES, start=2):
        swatch = ws.cell(row=r, column=4)
        swatch.fill = PatternFill("solid", fgColor="FF" + t["color"].lstrip("#"))
        swatch.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 58


def sheet_activities(wb):
    ws = wb.create_sheet("Activities")
    ws.append(["theme_no", "activity", "outcomes", "tags", "start_month",
               "end_month", "milestone_months", "links", "status", "notes"])
    for row in seed.ACTIVITIES:
        theme, act, out, tags, start, end, miles, links, notes = row
        ws.append([theme, act, out, tags, start, end, miles, links,
                   "Not started", notes])
    style_header(ws, 10, freeze="B2")
    set_widths(ws, [9, 52, 52, 22, 12, 11, 17, 46, 13, 38])
    body(ws, 2, 10, wrap_cols=(2, 3, 4, 8, 10))
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 46
        for col in (1, 5, 6, 7):
            ws.cell(row=r, column=col).alignment = Alignment(
                horizontal="center", vertical="top")

    last = ws.max_row + 300
    theme_dv = DataValidation(type="whole", operator="between", formula1=1,
                              formula2=99, allow_blank=True,
                              error="Use the theme_no from the Themes sheet.",
                              errorTitle="Unknown theme")
    month_dv = DataValidation(type="whole", operator="between", formula1=1,
                              formula2=600, allow_blank=True,
                              error="Project months are counted from 1.",
                              errorTitle="Not a project month")
    status_dv = DataValidation(type="list", allow_blank=True,
                               formula1='"' + ",".join(STATUSES) + '"')
    for dv in (theme_dv, month_dv, status_dv):
        ws.add_data_validation(dv)
    theme_dv.add(f"A2:A{last}")
    month_dv.add(f"E2:F{last}")
    status_dv.add(f"I2:I{last}")

    hint = ws.cell(row=1, column=4)
    hint.comment = None  # keep the file free of comment parts
    return ws


def sheet_budget(wb):
    ws = wb.create_sheet("Budget")
    ws.append(["theme_no", "item", "description", "amount_eur", "shared"])
    for row in seed.BUDGET:
        ws.append(list(row))
    style_header(ws, 5, freeze="B2")
    set_widths(ws, [9, 34, 74, 14, 10])
    body(ws, 2, 5, wrap_cols=(3,))
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=4).number_format = '€ #,##0'
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center",
                                                       vertical="top")
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center",
                                                       vertical="top")
    dv = DataValidation(type="list", allow_blank=True, formula1='"yes,no"')
    ws.add_data_validation(dv)
    dv.add(f"E2:E{ws.max_row + 200}")


def sheet_questions(wb):
    ws = wb.create_sheet("Open questions")
    ws.append(["theme_no", "question", "links"])
    for row in seed.OPEN_QUESTIONS:
        ws.append(list(row))
    style_header(ws, 3, freeze="B2")
    set_widths(ws, [9, 66, 60])
    body(ws, 2, 3, wrap_cols=(2, 3))
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 34


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    sheet_themes(wb)
    sheet_activities(wb)
    sheet_budget(wb)
    sheet_questions(wb)
    sheet_settings(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
